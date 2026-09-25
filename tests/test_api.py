"""Tests de la API REST. Responsable: Diana Dure (QA Tester)."""
import pytest


def test_health_ok(api_client):
    r = api_client.get("/health")
    assert r.status_code == 200
    assert r.json()["status"] == "ok"


def test_adaptar_flashcards_salud(api_client, documento_salud):
    body = {
        "documento_titulo": "Protocolo de Bioseguridad",
        "documento_contenido": documento_salud,
        "perfil_destinatario": "Principiante",
        "formato_salida": "Flashcards",
        "nicho_sector": "Salud",
    }
    r = api_client.post("/api/v1/adaptar", json=body)
    assert r.status_code == 200
    data = r.json()
    assert data["status"] == "exito"
    assert data["evaluacion_calidad"]["umbral_aplicado"] == 0.9
    assert data["almacenamiento_oci"]["status_upload"] in ("completado", "fallido_local")


@pytest.mark.parametrize(
    "formato,perfil",
    [
        ("Flashcards", "Principiante"),
        ("Quiz", "Desarrollador Junior/Semi Senior"),
        ("Tutorial", "Líder Técnico/Arquitecto"),
        ("Resumen Ejecutivo", "Gestor/Ejecutivo"),
        ("Guion de Clase", "Principiante"),
    ],
)
def test_adaptar_todos_los_formatos(api_client, documento_salud, formato, perfil):
    """Checklist del hackathon: probar >= 2 perfiles y >= 2 formatos. Aquí probamos los 5."""
    body = {
        "documento_titulo": "Protocolo de Bioseguridad",
        "documento_contenido": documento_salud,
        "perfil_destinatario": perfil,
        "formato_salida": formato,
        "nicho_sector": "Salud",
    }
    r = api_client.post("/api/v1/adaptar", json=body)
    assert r.status_code == 200, r.text
    assert r.json()["contenido_adaptado"]["formato"] == formato


def test_adaptar_valida_campos_faltantes(api_client):
    r = api_client.post("/api/v1/adaptar", json={"documento_titulo": "x"})
    assert r.status_code == 422
    assert r.json()["codigo"] == "VALIDACION_ENTRADA"


def test_adaptar_rechaza_perfil_no_soportado(api_client, documento_salud):
    body = {
        "documento_titulo": "Doc",
        "documento_contenido": documento_salud,
        "perfil_destinatario": "Doctorado en Medicina",  # no está en el enum
        "formato_salida": "Flashcards",
    }
    r = api_client.post("/api/v1/adaptar", json=body)
    assert r.status_code == 422


def test_adaptar_archivo_txt(api_client, documento_salud):
    files = {"archivo": ("protocolo.txt", documento_salud.encode("utf-8"), "text/plain")}
    data = {
        "perfil_destinatario": "Principiante",
        "formato_salida": "Flashcards",
        "nicho_sector": "Salud",
    }
    r = api_client.post("/api/v1/adaptar/archivo", files=files, data=data)
    assert r.status_code == 200, r.text


def test_adaptar_archivo_demasiado_grande_da_413(api_client):
    grande = b"x" * (11 * 1024 * 1024)
    files = {"archivo": ("grande.txt", grande, "text/plain")}
    data = {"perfil_destinatario": "Principiante", "formato_salida": "Flashcards"}
    r = api_client.post("/api/v1/adaptar/archivo", files=files, data=data)
    assert r.status_code == 413


def test_flujo_completo_guarda_y_recupera(api_client, documento_salud):
    body = {
        "documento_titulo": "Protocolo de Bioseguridad",
        "documento_contenido": documento_salud,
        "perfil_destinatario": "Principiante",
        "formato_salida": "Flashcards",
        "nicho_sector": "Salud",
    }
    r1 = api_client.post("/api/v1/adaptar", json=body)
    objeto_id = r1.json()["almacenamiento_oci"]["objeto_id"]

    r2 = api_client.get(f"/api/v1/contenidos/{objeto_id}")
    assert r2.status_code == 200
    assert r2.json()["request_id"] == r1.json()["request_id"]


def test_contenido_inexistente_da_404(api_client):
    r = api_client.get("/api/v1/contenidos/no-existe.json")
    assert r.status_code == 404


def test_front_end_se_sirve_en_la_raiz(api_client):
    r = api_client.get("/")
    assert r.status_code == 200
    assert "Kairos" in r.text
    assert api_client.get("/static/app.js").status_code == 200
    assert api_client.get("/static/styles.css").status_code == 200


def test_opciones_coinciden_con_los_enums(api_client):
    from nuevamente.schemas.enums import FormatoSalida, PerfilDestinatario

    op = api_client.get("/api/v1/opciones").json()
    assert op["perfiles"] == [p.value for p in PerfilDestinatario]
    assert op["formatos"] == [f.value for f in FormatoSalida]
    assert op["umbral_por_nicho"]["Salud"] == 0.9


def _crear(api_client, documento, formato):
    body = {
        "documento_titulo": "Protocolo de Bioseguridad",
        "documento_contenido": documento,
        "perfil_destinatario": "Principiante",
        "formato_salida": formato,
    }
    return api_client.post("/api/v1/adaptar", json=body).json()["almacenamiento_oci"]["objeto_id"]


def test_video_del_guion_de_clase_es_un_mp4(api_client, documento_salud):
    objeto_id = _crear(api_client, documento_salud, "Guion de Clase")
    r = api_client.get(f"/api/v1/contenidos/{objeto_id}/video", params={"titulo": "Protocolo"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "video/mp4"
    assert r.content[4:8] == b"ftyp"  # cabecera de un contenedor MP4


def test_video_solo_para_guion_de_clase(api_client, documento_salud):
    objeto_id = _crear(api_client, documento_salud, "Flashcards")
    assert api_client.get(f"/api/v1/contenidos/{objeto_id}/video").status_code == 422


def test_video_con_voz_masculina(api_client, documento_salud):
    objeto_id = _crear(api_client, documento_salud, "Guion de Clase")
    r = api_client.get(f"/api/v1/contenidos/{objeto_id}/video", params={"voz": "masculina"})
    assert r.status_code == 200, r.text
    assert r.content[4:8] == b"ftyp"


def test_video_rechaza_voz_desconocida(api_client, documento_salud):
    objeto_id = _crear(api_client, documento_salud, "Guion de Clase")
    assert api_client.get(f"/api/v1/contenidos/{objeto_id}/video", params={"voz": "robot"}).status_code == 422


def test_opciones_informan_voces_y_muestra_sin_voz_da_404(api_client):
    # en los tests VIDEO_TTS=off: no hay voces y la muestra lo dice claramente
    assert api_client.get("/api/v1/opciones").json()["voces"] == {"femenina": None, "masculina": None}
    assert api_client.get("/api/v1/voces/femenina/muestra").status_code == 404


@pytest.mark.parametrize("formato", ["Flashcards", "Quiz", "Tutorial", "Resumen Ejecutivo", "Guion de Clase"])
def test_exportar_presentacion_pptx(api_client, documento_salud, formato):
    import io

    from pptx import Presentation

    objeto_id = _crear(api_client, documento_salud, formato)
    r = api_client.get(f"/api/v1/contenidos/{objeto_id}/exportar", params={"formato": "pptx", "titulo": "Protocolo"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml")
    deck = Presentation(io.BytesIO(r.content))
    assert len(deck.slides) >= 3  # portada + contenido + cierre
    textos = " ".join(sh.text_frame.text for s in deck.slides for sh in s.shapes if sh.has_text_frame)
    assert "Protocolo" in textos and "¡Gracias!" in textos
    if formato == "Guion de Clase":  # la narración va en las notas del orador
        assert deck.slides[1].notes_slide.notes_text_frame.text.startswith("Narración")

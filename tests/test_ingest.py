"""Tests de ingesta y chunking. Responsable: Diana Dure (QA Tester)."""
import pytest

from nuevamente.ingest.chunking import chunkear_documento
from nuevamente.ingest.readers import DocumentoInvalidoError, leer_texto_plano


def test_chunking_respeta_secciones(documento_salud):
    chunks = chunkear_documento("Protocolo", documento_salud)
    secciones = {c.seccion for c in chunks}
    assert "Lavado de manos" in secciones
    assert "Equipos de proteccion personal" in secciones
    assert "Manejo de residuos" in secciones


def test_chunking_genera_chunk_ids_unicos(documento_salud):
    chunks = chunkear_documento("Protocolo", documento_salud)
    ids = [c.chunk_id for c in chunks]
    assert len(ids) == len(set(ids))


def test_chunking_documento_corto_da_un_solo_chunk():
    chunks = chunkear_documento("Doc corto", "Una sola oración corta de prueba.")
    assert len(chunks) == 1


def test_leer_texto_plano_rechaza_vacio():
    with pytest.raises(DocumentoInvalidoError):
        leer_texto_plano("Doc", "   ")


def test_leer_texto_plano_ok():
    doc = leer_texto_plano("Doc", "Contenido real del documento.")
    assert doc.titulo == "Doc"
    assert "real" in doc.contenido


def _pdf_minimo(dibujos) -> bytes:
    """PDF de una página con los textos (x, y, texto) en el orden dado dentro del archivo."""
    flujo = "\n".join(f"BT /F1 11 Tf {x} {y} Td ({t}) Tj ET" for x, y, t in dibujos).encode("latin-1")
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(flujo) + flujo + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    salida, posiciones = b"%PDF-1.4\n", []
    for i, obj in enumerate(objetos, start=1):
        posiciones.append(len(salida))
        salida += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(salida)
    salida += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objetos) + 1)
    salida += b"".join(b"%010d 00000 n \n" % p for p in posiciones)
    salida += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objetos) + 1, xref)
    return salida


def test_pdf_con_titulos_guardados_antes_que_el_texto(tmp_path):
    """El PDF trae primero los dos títulos y el pie, y luego el texto: cada título debe
    quedar con su propio contenido, no todo bajo el último."""
    from nuevamente.ingest.readers import leer_pdf

    ruta = tmp_path / "manual.pdf"
    ruta.write_bytes(_pdf_minimo([
        (72, 700, "1 Obligaciones del paciente"),
        (72, 600, "2 Regimen de visitas"),
        (72, 60, "Centro Medico General Pagina 1 de 3"),
        (72, 680, "Trato respetuoso: se requiere mantener un trato cortes con el personal."),
        (72, 580, "Horario de visitas: las visitas se permiten de 10:00 a 12:00 horas."),
    ]))
    chunks = chunkear_documento("Manual", leer_pdf(ruta).contenido)
    assert [(c.seccion, c.texto.split(":")[0]) for c in chunks] == [
        ("Obligaciones del paciente", "Trato respetuoso"),
        ("Regimen de visitas", "Horario de visitas"),
    ]
    assert not any("Pagina 1 de 3" in c.texto for c in chunks)


def test_titulos_apilados_no_dejan_el_texto_bajo_un_titulo_ajeno():
    """Si no se puede saber de qué título es cada parte, la sección lleva los dos."""
    texto = (
        "MANUAL DEL PACIENTE\nNormas de la institucion para pacientes y familiares.\n"
        "1 Obligaciones del paciente\n2 Regimen de visitas\n"
        "Trato respetuoso: se requiere un trato cortes.\nHorario de visitas: de 10:00 a 12:00 horas.\n"
    )
    secciones = [c.seccion for c in chunkear_documento("Manual", texto)]
    assert secciones == ["Manual del paciente", "Obligaciones del paciente · Regimen de visitas"]


def test_indice_no_se_toma_por_titulos_apilados():
    texto = (
        "1 Acceso\n2 Modos\n"
        "1 Acceso\nSe accede por la consola del equipo.\n2 Modos\nHay modo usuario y modo privilegiado.\n"
    )
    assert [c.seccion for c in chunkear_documento("Manual", texto)] == ["Acceso", "Modos"]


def test_texto_pegado_sin_saltos_de_linea_recupera_sus_apartados():
    texto = (
        "Protocolo de BioseguridadUn protocolo de bioseguridad es un conjunto de normas preventivas."
        "Te presento una guía adaptada a los sectores más comunes:"
        "🏥 1. Ámbito de la SaludEn consultorios las medidas se dividen por niveles de contención."
        "Equipo de Protección Personal (EPP):Uso obligatorio de batas y mascarillas."
        "🏢 2. Ámbito LaboralOrientado a mantener espacios de trabajo seguros.Higiene de manos con agua y jabón."
    )
    chunks = chunkear_documento("Protocolo", texto)
    assert [c.seccion for c in chunks] == ["Introducción", "Ámbito de la Salud", "Ámbito Laboral"]
    assert chunks[1].texto.startswith("En consultorios") and "batas y mascarillas" in chunks[1].texto
    assert "Higiene de manos" in chunks[2].texto and "contención" not in chunks[2].texto


def test_documento_normal_no_se_altera():
    from nuevamente.ingest.readers import limpiar_texto

    texto = "Usamos JavaScript y PowerPoint. La versión 2.0 salió ayer.\nSegunda línea: sin cambios."
    assert limpiar_texto(texto) == texto


def test_limpieza_quita_indice_avisos_legales_y_contacto():
    from nuevamente.ingest.readers import limpiar_texto

    texto = (
        "Algoritmo ........ 4\nDatos …… 12\n"
        "© 2023 Universidad Nacional\nTodos los derechos reservados.\nISBN 978-3-16-148410-0\n"
        "Esta obra está bajo una licencia Creative Commons.\n"
        "Contacto: info@universidad.edu\nsoporte@universidad.edu\n"
        "Algoritmo: fórmula o conjunto de reglas para resolver un problema.\n"
        "Código abierto: software cuya licencia permite usarlo y modificarlo.\n"
        "Ingresa a https://cloud.oracle.com\n"
    )
    assert limpiar_texto(texto) == (
        "Algoritmo: fórmula o conjunto de reglas para resolver un problema.\n"
        "Código abierto: software cuya licencia permite usarlo y modificarlo.\n"
        "Ingresa a https://cloud.oracle.com"
    )


def test_pdf_quita_encabezados_pies_y_numeros_de_pagina():
    from nuevamente.ingest.readers import _sin_bordes_repetidos

    paginas = [
        ["Manual de IA · DD-IA-01", f"Texto propio de la página {n}.", "Otra línea del cuerpo.", "Universidad Nacional", str(n)]
        for n in range(1, 5)
    ]
    assert _sin_bordes_repetidos(paginas) == [
        [f"Texto propio de la página {n}.", "Otra línea del cuerpo."] for n in range(1, 5)
    ]


def test_pdf_corto_conserva_lineas_que_no_se_repiten():
    from nuevamente.ingest.readers import _sin_bordes_repetidos

    paginas = [["Introducción", "Primera idea del texto.", "3"], ["Conclusión", "Última idea."]]
    assert _sin_bordes_repetidos(paginas) == [
        ["Introducción", "Primera idea del texto."],
        ["Conclusión", "Última idea."],
    ]


def test_conceptos_clave_de_un_texto_sin_titulos():
    from nuevamente.ingest.conceptos import extraer_conceptos

    texto = (
        "La Virtual Cloud Network (VCN) es una red privada configurada en Oracle Cloud Infrastructure. "
        "Similar a una red tradicional, incluye Internet Gateways, NAT Gateways y Security Lists. "
        "La Organización Mundial de la Salud publica guías."
    )
    assert extraer_conceptos(texto) == [
        "VCN", "Oracle Cloud Infrastructure", "Internet Gateways", "NAT Gateways", "Security Lists",
        "Organización Mundial de la Salud",
    ]


def _docx_minimo(ruta):
    from docx import Document

    doc = Document()
    doc.add_heading("Protocolo de visitas", level=1)
    doc.add_paragraph("Las visitas se permiten de 10:00 a 12:00 horas en todas las salas.")
    doc.add_heading("Equipo de protección", level=2)
    doc.add_paragraph("El personal debe usar guantes y mascarilla durante la atención.")
    tabla = doc.add_table(rows=2, cols=2)
    tabla.cell(0, 0).text, tabla.cell(0, 1).text = "Elemento", "Uso"
    tabla.cell(1, 0).text, tabla.cell(1, 1).text = "Guantes", "Contacto con fluidos"
    doc.save(str(ruta))


def test_docx_conserva_titulos_y_tablas(tmp_path):
    """Los títulos de Word separan las secciones y las tablas se leen fila a fila."""
    from nuevamente.ingest.readers import leer_documento

    ruta = tmp_path / "protocolo.docx"
    _docx_minimo(ruta)
    leido = leer_documento(ruta)
    assert leido.titulo == "protocolo"
    assert "Guantes | Contacto con fluidos" in leido.contenido
    secciones = [c.seccion for c in chunkear_documento("Protocolo", leido.contenido)]
    assert "Protocolo de visitas" in secciones and "Equipo de protección" in secciones


def test_docx_danado_es_documento_invalido(tmp_path):
    from nuevamente.ingest.readers import leer_documento

    ruta = tmp_path / "roto.docx"
    ruta.write_bytes(b"esto no es un docx")
    with pytest.raises(DocumentoInvalidoError):
        leer_documento(ruta)


# ---------- Títulos de sección (requisitos del chunking) ----------

_CUERPO = "Este párrafo explica el tema con suficiente detalle para formar un fragmento de texto corrido."


def _secciones(texto):
    vistas = []
    for c in chunkear_documento("Doc", texto):
        if c.seccion not in vistas:
            vistas.append(c.seccion)
    return vistas


def test_reconoce_encabezados_markdown_de_tres_niveles():
    md = f"# Redes en la nube\n{_CUERPO}\n\n## Configuración de IP\n{_CUERPO}\n\n### Rutas OSPF\n{_CUERPO}\n"
    assert _secciones(md) == ["Redes en la nube", "Configuración de IP", "Rutas OSPF"]


def test_en_texto_plano_reconoce_titulos_numerados_y_en_mayusculas():
    txt = (f"1. Acceso a la consola\n{_CUERPO}\n\n2. CONFIGURACIÓN DE DIRECCIONES IP\n{_CUERPO}\n\n"
           f"MANEJO DE RUTAS OSPF\n{_CUERPO}\n")
    assert _secciones(txt) == ["Acceso a la consola", "Configuración de direcciones IP", "Manejo de rutas OSPF"]


def test_titulos_sin_numeracion_en_tipo_oracion_y_con_siglas():
    """Las palabras comunes cortas en mayúsculas (GUÍA, USO, RED) no se toman por siglas."""
    md = (f"# 1. GUÍA DE OSPF Y BGP\n{_CUERPO} Esta guía usa la red.\n\n"
          f"# 2.1 USO DEL EPP EN LA RED\n{_CUERPO}\n\n# 3. INTRODUCCIÓN A TCP/IP E IPv4\n{_CUERPO}\n")
    assert _secciones(md) == ["Guía de OSPF y BGP", "Uso del EPP en la red", "Introducción a TCP/IP e IPv4"]


def test_una_sigla_del_documento_se_conserva_aunque_tenga_vocales():
    """KAIROS no es una sigla conocida, pero el documento la escribe así en el texto."""
    md = f"# MANUAL DE KAIROS\nEl sistema KAIROS adapta documentos para cada perfil del equipo.\n\n# OTRO APARTADO\n{_CUERPO}\n"
    assert _secciones(md)[0] == "Manual de KAIROS"


def test_el_texto_antes_del_primer_titulo_es_la_introduccion():
    md = f"Texto previo al primer título que presenta el documento completo.\n\n# Primer tema\n{_CUERPO}\n"
    assert _secciones(md) == ["Introducción", "Primer tema"]
    assert _secciones(f"# Primer tema\n{_CUERPO}\n") == ["Primer tema"]  # sin texto previo, sin introducción


def test_items_de_lista_y_oraciones_partidas_no_son_titulos():
    """En un PDF, "1. Rótulo: La oración…" es un punto de lista, y una línea que sigue en
    minúscula en la de abajo es una oración partida: ninguno abre una sección."""
    txt = (f"1. Acceso a la consola\n{_CUERPO}\n"
           "1. Enfoque Cognitivo: Este marco resulta idóneo para comprender la\nrigidez del paciente en su entorno.\n"
           "2. Interacción y Colapso: La combinación de un\nNeuroticismo alto con baja Amabilidad.\n\n"
           f"2. Configuración de la red\n{_CUERPO}\n")
    assert _secciones(txt) == ["Acceso a la consola", "Configuración de la red"]


def test_titulo_en_dos_lineas_y_nombres_propios():
    txt = (f"EVALUACIÓN INTEGRAL EN CASO\nCOMPLEJO (PACIENTE RODRIGO ESPINOZA)\n\nEl informe analiza el caso de Rodrigo Espinoza con detalle suficiente.\n\n"
           f"MANIFESTACIONES CLÍNICAS DEL CASO\n{_CUERPO}\n")
    assert _secciones(txt)[0] == "Evaluación integral en caso complejo (paciente Rodrigo Espinoza)"


# ---------- OCR de PDF escaneados (HU-08) ----------

def _pdf_escaneado(tmp_path):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.rect(20, 20, 100, 60)  # solo dibujo, ningún texto: como una página escaneada
    ruta = tmp_path / "escaneo.pdf"
    pdf.output(str(ruta))
    return ruta


def _con_gemini(monkeypatch, texto_ocr):
    import dataclasses
    from types import SimpleNamespace

    from google import genai

    from nuevamente import config

    monkeypatch.setattr(config, "settings", dataclasses.replace(config.settings, gemini_api_key="clave-de-prueba", ocr="auto"))
    pedidos = []

    class _Modelos:
        def generate_content(self, model, contents, config=None):
            pedidos.append((model, contents[0].inline_data.mime_type))
            return SimpleNamespace(text=texto_ocr)

    monkeypatch.setattr(genai, "Client", lambda **_: SimpleNamespace(models=_Modelos()))
    return pedidos


def test_pdf_escaneado_se_lee_con_ocr(monkeypatch, tmp_path):
    from nuevamente.ingest.readers import leer_pdf

    texto = "1. Lavado de manos\nEl lavado de manos con agua y jabón elimina la mayoría de los microorganismos."
    pedidos = _con_gemini(monkeypatch, texto)
    leido = leer_pdf(_pdf_escaneado(tmp_path))
    assert "elimina la mayoría de los microorganismos" in leido.contenido
    assert pedidos[0][1] == "application/pdf"  # Gemini recibe el PDF tal cual


def test_ocr_sin_texto_suficiente_avisa(monkeypatch, tmp_path):
    from nuevamente.ingest.readers import leer_pdf

    _con_gemini(monkeypatch, "  ")
    with pytest.raises(DocumentoInvalidoError, match="no reconoció texto suficiente"):
        leer_pdf(_pdf_escaneado(tmp_path))


def test_sin_api_key_el_pdf_escaneado_explica_que_falta(monkeypatch, tmp_path):
    import dataclasses

    from nuevamente import config
    from nuevamente.ingest.readers import leer_pdf

    monkeypatch.setattr(config, "settings", dataclasses.replace(config.settings, gemini_api_key="", gemini_api_keys=()))
    with pytest.raises(DocumentoInvalidoError, match="no tiene texto extraíble"):
        leer_pdf(_pdf_escaneado(tmp_path))

"""Presentación PowerPoint (.pptx) a partir de cualquier contenido adaptado.

Responsable en el equipo Kairos G10: Dario Higuera Moreno (No Code Developer).

Estructura: portada → diapositivas del formato → cierre. Cada formato tiene su
propio diseño (tarjetas, preguntas con su respuesta, pasos, puntos clave,
escenas) y las notas del orador llevan lo que el presentador necesita decir:
la respuesta de cada pregunta o la narración de cada escena del guion.

Paleta y tipografía siguen la interfaz web; se usa Arial porque viene con Office
y se ve igual en PowerPoint, Keynote y LibreOffice.
"""
from __future__ import annotations

import io
import math
from functools import lru_cache
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from nuevamente.schemas.formatos import (
    ContenidoAdaptado,
    FlashcardsContenido,
    GuionDeClaseContenido,
    QuizContenido,
    ResumenEjecutivoContenido,
    TutorialContenido,
)

ANCHO_IN, ALTO_IN = 13.333, 7.5
MARGEN_IN = 0.7
FUENTE = "Arial"

PRIMARIO = RGBColor(0x6C, 0x4E, 0xF5)
PRIMARIO_SUAVE = RGBColor(0xEF, 0xEA, 0xFF)
CORAL = RGBColor(0xFF, 0x7A, 0x59)
CORAL_SUAVE = RGBColor(0xFF, 0xE9, 0xE2)
MENTA = RGBColor(0x12, 0xA3, 0x7F)
MENTA_SUAVE = RGBColor(0xDC, 0xF7, 0xEE)
SOL_SUAVE = RGBColor(0xFF, 0xF4, 0xD4)
TEXTO = RGBColor(0x2D, 0x24, 0x40)
TEXTO_SUAVE = RGBColor(0x6E, 0x64, 0x80)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

_FUENTES_MEDIDA = {
    False: [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",  # mismas métricas que Arial
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",  # más ancha: estimación conservadora
    ],
    True: [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "/usr/share/fonts/truetype/msttcorefonts/Arial_Bold.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    ],
}
_INTERLINEADO = 1.22


# ---------- Medición de texto ----------

@lru_cache(maxsize=128)
def _fuente_medida(pt: int, negrita: bool):
    for ruta in _FUENTES_MEDIDA[negrita]:
        if Path(ruta).exists():
            return ImageFont.truetype(ruta, pt)
    return None


def _lineas(texto: str, pt: int, ancho_in: float, negrita: bool) -> int:
    """Cuántas líneas ocupa `texto` a `pt` puntos en una caja de `ancho_in` pulgadas."""
    ancho_pt = ancho_in * 72 * 0.97  # pequeño margen de seguridad frente al motor de cada programa
    fuente = _fuente_medida(pt, negrita)
    medir = fuente.getlength if fuente else (lambda s: len(s) * pt * (0.58 if negrita else 0.52))
    total = 0
    for parrafo in texto.split("\n"):
        lineas, actual = 1, ""
        for palabra in parrafo.split():
            if medir(palabra) > ancho_pt:
                return 10_000  # una palabra más ancha que la caja se partiría: este tamaño no sirve
            prueba = f"{actual} {palabra}".strip()
            if medir(prueba) <= ancho_pt or not actual:
                actual = prueba
            else:
                lineas += 1
                actual = palabra
        total += lineas
    return total


def _alto(texto: str, pt: int, ancho_in: float, negrita=False) -> float:
    """Alto en pulgadas que ocupa `texto` a `pt` puntos."""
    return _lineas(texto, pt, ancho_in, negrita) * pt * _INTERLINEADO / 72


def _tamano(texto: str, ancho_in: float, alto_in: float, maximo: int, minimo: int, negrita=False) -> int:
    """Mayor tamaño (pt) con el que el texto cabe en la caja; el mínimo si ninguno cabe."""
    for pt in range(maximo, minimo - 1, -1):
        if _alto(texto, pt, ancho_in, negrita) <= alto_in:
            return pt
    return minimo


def _recortar(texto: str, ancho_in: float, alto_in: float, pt: int, negrita=False) -> str:
    """Si ni al tamaño mínimo cabe, recorta por palabras y agrega puntos suspensivos."""
    if _alto(texto, pt, ancho_in, negrita) <= alto_in:
        return texto
    palabras = texto.split()
    while palabras and _alto(" ".join(palabras) + "…", pt, ancho_in, negrita) > alto_in:
        palabras.pop()
    return " ".join(palabras).rstrip(".,;:") + "…"


# ---------- Primitivas ----------

class _Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(ANCHO_IN)
        self.prs.slide_height = Inches(ALTO_IN)
        self._vacio = self.prs.slide_layouts[6]

    def diapositiva(self, fondo: str = "blanco"):
        slide = self.prs.slides.add_slide(self._vacio)
        relleno = slide.background.fill
        if fondo == "degradado":
            relleno.gradient()
            relleno.gradient_angle = 0
            relleno.gradient_stops[0].color.rgb = PRIMARIO
            relleno.gradient_stops[1].color.rgb = CORAL
        else:
            relleno.solid()
            relleno.fore_color.rgb = BLANCO
        return slide

    def guardar(self) -> bytes:
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()


def _texto(slide, x, y, w, h, texto, pt, color=TEXTO, negrita=False, alinear=PP_ALIGN.LEFT,
           ancla=MSO_ANCHOR.TOP, ajustar=True, minimo=12):
    """Caja de texto sin márgenes internos; si `ajustar`, reduce la letra hasta que quepa."""
    if ajustar:
        pt = _tamano(texto, w, h, pt, minimo, negrita)
        texto = _recortar(texto, w, h, pt, negrita)
    caja = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    marco = caja.text_frame
    marco.word_wrap = True
    marco.auto_size = MSO_AUTO_SIZE.NONE
    marco.margin_left = marco.margin_right = marco.margin_top = marco.margin_bottom = Emu(0)
    marco.vertical_anchor = ancla
    for i, linea in enumerate(texto.split("\n")):
        parrafo = marco.paragraphs[0] if i == 0 else marco.add_paragraph()
        parrafo.alignment = alinear
        run = parrafo.add_run()
        run.text = linea
        run.font.name = FUENTE
        run.font.size = Pt(pt)
        run.font.bold = negrita
        run.font.color.rgb = color
    return caja


def _caja(slide, x, y, w, h, relleno, forma=MSO_SHAPE.ROUNDED_RECTANGLE, radio=0.08):
    figura = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    figura.fill.solid()
    figura.fill.fore_color.rgb = relleno
    figura.line.fill.background()
    figura.shadow.inherit = False
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        figura.adjustments[0] = radio
    return figura


def _circulo(slide, x, y, d, relleno, texto, pt, color=BLANCO):
    _caja(slide, x, y, d, d, relleno, forma=MSO_SHAPE.OVAL)
    _texto(slide, x, y, d, d, texto, pt, color, negrita=True, alinear=PP_ALIGN.CENTER,
           ancla=MSO_ANCHOR.MIDDLE, ajustar=False)


def _encabezado(slide, etiqueta: str, titulo: str, alto_titulo=1.1):
    """Etiqueta pequeña en coral + título grande. Devuelve la `y` donde empieza el cuerpo."""
    _texto(slide, MARGEN_IN, 0.55, ANCHO_IN - 2 * MARGEN_IN, 0.35, etiqueta.upper(), 13, CORAL, negrita=True,
           ajustar=False)
    _texto(slide, MARGEN_IN, 0.95, ANCHO_IN - 2 * MARGEN_IN, alto_titulo, titulo, 34, TEXTO, negrita=True,
           ancla=MSO_ANCHOR.TOP, minimo=22)
    return 0.95 + alto_titulo + 0.35


def _columnas(slide, y: float, bloques: list[tuple[str, str, RGBColor]], maximo=22, minimo=12):
    """Tarjetas lado a lado con título y texto; el alto se ajusta al contenido más largo."""
    ancho = (ANCHO_IN - 2 * MARGEN_IN - 0.4 * (len(bloques) - 1)) / len(bloques)
    disponible = ALTO_IN - y - 0.8
    pt = min(_tamano(texto, ancho - 0.8, disponible - 1.25, maximo, minimo) for _, texto, _ in bloques)
    alto = min(disponible, max(_alto(texto, pt, ancho - 0.8) for _, texto, _ in bloques) + 1.3)
    for k, (titulo, texto, color) in enumerate(bloques):
        x = MARGEN_IN + k * (ancho + 0.4)
        _caja(slide, x, y, ancho, alto, color)
        _texto(slide, x + 0.4, y + 0.3, ancho - 0.8, 0.45, titulo, 20, TEXTO, negrita=True, ajustar=False)
        _texto(slide, x + 0.4, y + 0.95, ancho - 0.8, alto - 1.25, texto, pt, TEXTO, ajustar=False)


def _pie(slide, texto: str = "Kairos"):
    _texto(slide, MARGEN_IN, ALTO_IN - 0.55, 4, 0.3, texto, 11, TEXTO_SUAVE, negrita=True, ajustar=False)


def _notas(slide, texto: str):
    slide.notes_slide.notes_text_frame.text = texto


# ---------- Portada y cierre ----------

def _portada(deck: _Deck, etiqueta: str, titulo: str, subtitulo: str):
    slide = deck.diapositiva("degradado")
    _texto(slide, MARGEN_IN + 0.1, 1.6, 10, 0.4, etiqueta.upper(), 16, BLANCO, negrita=True, ajustar=False)
    _texto(slide, MARGEN_IN + 0.1, 2.1, ANCHO_IN - 2 * MARGEN_IN - 0.2, 2.2, titulo, 48, BLANCO, negrita=True,
           ancla=MSO_ANCHOR.TOP, minimo=30)
    _texto(slide, MARGEN_IN + 0.1, 4.5, 10.5, 1.3, subtitulo, 20, BLANCO, minimo=14)
    _texto(slide, MARGEN_IN + 0.1, ALTO_IN - 0.95, 4, 0.45, "Kairos", 20, BLANCO, negrita=True, ajustar=False)


def _cierre(deck: _Deck, lineas: list[str]):
    slide = deck.diapositiva("degradado")
    _texto(slide, MARGEN_IN + 0.1, 2.3, 10, 1.3, "¡Gracias!", 66, BLANCO, negrita=True, ajustar=False)
    _texto(slide, MARGEN_IN + 0.1, 3.8, 11, 1.6, "\n".join(lineas), 20, BLANCO, minimo=14)
    _texto(slide, MARGEN_IN + 0.1, ALTO_IN - 0.95, 4, 0.45, "Kairos", 20, BLANCO, negrita=True, ajustar=False)


# ---------- Formatos ----------

def _flashcards(deck: _Deck, c: FlashcardsContenido):
    total = len(c.items)
    for i, item in enumerate(c.items, 1):
        slide = deck.diapositiva()
        _texto(slide, MARGEN_IN, 0.55, 8, 0.35, f"TARJETA {i} DE {total}", 13, CORAL, negrita=True, ajustar=False)
        # frente: tarjeta violeta a la izquierda
        _caja(slide, MARGEN_IN, 1.2, 5.3, 5.2, PRIMARIO)
        _texto(slide, MARGEN_IN + 0.45, 1.6, 4.4, 0.35, "PREGUNTA", 13, PRIMARIO_SUAVE, negrita=True, ajustar=False)
        _texto(slide, MARGEN_IN + 0.45, 2.1, 4.4, 3.9, item.frente, 30, BLANCO, negrita=True, minimo=18)
        # dorso: respuesta y pista a la derecha
        x = MARGEN_IN + 5.3 + 0.4
        ancho = ANCHO_IN - MARGEN_IN - x
        alto_resp = 3.55 if item.pista_didactica else 5.2
        _caja(slide, x, 1.2, ancho, alto_resp, PRIMARIO_SUAVE)
        _texto(slide, x + 0.45, 1.6, ancho - 0.9, 0.35, "RESPUESTA", 13, PRIMARIO, negrita=True, ajustar=False)
        _texto(slide, x + 0.45, 2.1, ancho - 0.9, alto_resp - 1.3, item.dorso, 22, TEXTO, minimo=13)
        if item.pista_didactica:
            _caja(slide, x, 4.95, ancho, 1.45, SOL_SUAVE)
            _texto(slide, x + 0.45, 5.15, ancho - 0.9, 0.3, "PISTA", 12, TEXTO_SUAVE, negrita=True, ajustar=False)
            _texto(slide, x + 0.45, 5.5, ancho - 0.9, 0.75, item.pista_didactica, 16, TEXTO, minimo=11)
        _pie(slide)
        _notas(slide, f"Pregunta: {item.frente}\nRespuesta: {item.dorso}")


def _opciones(slide, pregunta, y0: float):
    """Las 4 opciones en una cuadrícula 2×2, con su letra en un círculo."""
    ancho = (ANCHO_IN - 2 * MARGEN_IN - 0.4) / 2
    alto = (ALTO_IN - y0 - 0.8 - 0.3) / 2
    textos = [o for o in pregunta.opciones]
    pt = min(_tamano(t, ancho - 1.4, alto - 0.3, 18, 11) for t in textos)  # mismo tamaño en las 4
    for j, opcion in enumerate(textos):
        fila, col = divmod(j, 2)
        x = MARGEN_IN + col * (ancho + 0.4)
        y = y0 + fila * (alto + 0.3)
        _caja(slide, x, y, ancho, alto, PRIMARIO_SUAVE)
        _circulo(slide, x + 0.3, y + (alto - 0.6) / 2, 0.6, PRIMARIO, chr(65 + j), 18)
        _texto(slide, x + 1.15, y + 0.15, ancho - 1.4, alto - 0.3, opcion, pt, TEXTO, ancla=MSO_ANCHOR.MIDDLE,
               ajustar=False)


def _quiz(deck: _Deck, c: QuizContenido):
    total = len(c.preguntas)
    for i, p in enumerate(c.preguntas, 1):
        letra = chr(65 + p.indice_correcto)
        opcion = p.opciones[p.indice_correcto]
        # 1) la pregunta
        slide = deck.diapositiva()
        y = _encabezado(slide, f"Pregunta {i} de {total}", p.enunciado, alto_titulo=1.2)
        _opciones(slide, p, y)
        _pie(slide)
        _notas(slide, f"Respuesta correcta: {letra}. {opcion}\n{p.justificacion}")
        # 2) la respuesta: solo la correcta, en grande, y su justificación si aporta algo
        slide = deck.diapositiva()
        y = _encabezado(slide, f"Respuesta · pregunta {i} de {total}", p.enunciado, alto_titulo=1.2)
        ancho = ANCHO_IN - 2 * MARGEN_IN
        aporta = opcion.strip(" .") not in p.justificacion
        alto = 2.4 if aporta else ALTO_IN - y - 0.8
        _caja(slide, MARGEN_IN, y, ancho, alto, MENTA_SUAVE)
        _circulo(slide, MARGEN_IN + 0.45, y + (alto - 1.0) / 2, 1.0, MENTA, "✓", 32)
        _texto(slide, MARGEN_IN + 1.85, y + 0.3, ancho - 2.3, 0.35, f"RESPUESTA CORRECTA: {letra}", 13, MENTA,
               negrita=True, ajustar=False)
        _texto(slide, MARGEN_IN + 1.85, y + 0.75, ancho - 2.3, alto - 1.05, opcion, 26, TEXTO, negrita=True,
               ancla=MSO_ANCHOR.MIDDLE, minimo=14)
        if aporta:
            y2 = y + alto + 0.35
            _texto(slide, MARGEN_IN, y2, ancho, 0.35, "POR QUÉ", 13, TEXTO_SUAVE, negrita=True, ajustar=False)
            _texto(slide, MARGEN_IN, y2 + 0.45, ancho, ALTO_IN - y2 - 1.25, p.justificacion, 18, TEXTO, minimo=11)
        _pie(slide)
        _notas(slide, f"Respuesta correcta: {letra}. {opcion}\n{p.justificacion}")


def _tutorial(deck: _Deck, c: TutorialContenido):
    slide = deck.diapositiva()
    y = _encabezado(slide, "Tutorial", "Objetivo")
    _caja(slide, MARGEN_IN, y, ANCHO_IN - 2 * MARGEN_IN, 2.2, PRIMARIO_SUAVE)
    _texto(slide, MARGEN_IN + 0.5, y + 0.35, ANCHO_IN - 2 * MARGEN_IN - 1, 1.5, c.objetivo, 26, TEXTO, negrita=True,
           ancla=MSO_ANCHOR.MIDDLE, minimo=16)
    if c.prerrequisitos:
        _texto(slide, MARGEN_IN, y + 2.6, ANCHO_IN - 2 * MARGEN_IN, 1.5,
               "Antes de empezar: " + " · ".join(c.prerrequisitos), 16, TEXTO_SUAVE, minimo=11)
    _pie(slide)

    total = len(c.pasos)
    for paso in c.pasos:
        slide = deck.diapositiva()
        _texto(slide, MARGEN_IN, 0.55, 8, 0.35, f"PASO {paso.orden} DE {total}", 13, CORAL, negrita=True,
               ajustar=False)
        _circulo(slide, MARGEN_IN, 1.2, 1.5, CORAL, str(paso.orden), 44)
        x = MARGEN_IN + 1.5 + 0.5
        ancho = ANCHO_IN - MARGEN_IN - x
        _texto(slide, x, 1.25, ancho, 1.2, paso.titulo, 32, TEXTO, negrita=True, minimo=20)
        limite = ALTO_IN - 0.8
        alto_instr = 2.6 if paso.resultado_esperado else limite - 2.65
        pt = _tamano(paso.instruccion, ancho, alto_instr, 22, 14)
        _texto(slide, x, 2.65, ancho, alto_instr, paso.instruccion, pt, TEXTO, ajustar=False)
        if paso.resultado_esperado:
            y_res = 2.65 + min(alto_instr, _alto(paso.instruccion, pt, ancho)) + 0.45
            resultado = "Resultado esperado: " + paso.resultado_esperado
            pt_res = _tamano(resultado, ancho - 0.8, limite - y_res - 0.5, 18, 12)
            alto_res = min(limite - y_res, _alto(resultado, pt_res, ancho - 0.8) + 0.5)
            _caja(slide, x, y_res, ancho, alto_res, MENTA_SUAVE)
            _texto(slide, x + 0.4, y_res + 0.25, ancho - 0.8, alto_res - 0.5, resultado, pt_res, TEXTO,
                   ancla=MSO_ANCHOR.MIDDLE, minimo=11)
        _pie(slide)
        _notas(slide, paso.instruccion)

    if c.errores_comunes or c.checklist_final:
        slide = deck.diapositiva()
        y = _encabezado(slide, "Tutorial", "Antes de terminar")
        bloques = [(t, "\n".join(marca + i for i in items), color) for t, items, marca, color in (
            ("Errores comunes", c.errores_comunes, "• ", CORAL_SUAVE),
            ("Checklist final", c.checklist_final, "☐ ", MENTA_SUAVE),
        ) if items]
        _columnas(slide, y, bloques, maximo=20)
        _pie(slide)


def _resumen(deck: _Deck, c: ResumenEjecutivoContenido):
    slide = deck.diapositiva()
    y = _encabezado(slide, "Resumen ejecutivo", "En pocas palabras")
    _texto(slide, MARGEN_IN, y, ANCHO_IN - 2 * MARGEN_IN, ALTO_IN - y - 0.8, c.resumen, 22, TEXTO, minimo=13)
    _pie(slide)
    _notas(slide, c.resumen)

    slide = deck.diapositiva()
    y = _encabezado(slide, "Resumen ejecutivo", "Puntos clave")
    n = len(c.puntos_clave)
    columnas = 2 if n > 1 else 1
    filas = math.ceil(n / columnas)
    ancho = (ANCHO_IN - 2 * MARGEN_IN - 0.35 * (columnas - 1)) / columnas
    alto = (ALTO_IN - y - 0.8 - 0.35 * (filas - 1)) / filas
    pt = min(_tamano(p, ancho - 1.55, alto - 0.5, 24, 13, negrita=True) for p in c.puntos_clave)
    for k, punto in enumerate(c.puntos_clave):
        fila, col = divmod(k, columnas)
        x = MARGEN_IN + col * (ancho + 0.35)
        yy = y + fila * (alto + 0.35)
        _caja(slide, x, yy, ancho, alto, PRIMARIO_SUAVE)
        _circulo(slide, x + 0.35, yy + (alto - 0.7) / 2, 0.7, PRIMARIO, str(k + 1), 20)
        _texto(slide, x + 1.3, yy + 0.25, ancho - 1.55, alto - 0.5, punto, pt, TEXTO, negrita=True,
               ancla=MSO_ANCHOR.MIDDLE, ajustar=False)
    _pie(slide)

    if c.decisiones_o_riesgos or c.impacto_de_negocio:
        slide = deck.diapositiva()
        y = _encabezado(slide, "Resumen ejecutivo", "Riesgos e impacto")
        bloques = [(t, txt, color) for t, txt, color in (
            ("Riesgos y decisiones", "\n".join("• " + r for r in c.decisiones_o_riesgos), CORAL_SUAVE),
            ("Impacto de negocio", c.impacto_de_negocio, MENTA_SUAVE),
        ) if txt]
        _columnas(slide, y, bloques, maximo=24)
        _pie(slide)


def _reloj(segundos: int) -> str:
    return f"{segundos // 60}:{segundos % 60:02d}"


def _guion(deck: _Deck, c: GuionDeClaseContenido):
    total = len(c.escenas)
    inicio = 0
    for e in c.escenas:
        fin = inicio + e.duracion_seg
        slide = deck.diapositiva()
        titulo = e.apoyo_visual.split(":", 1)[-1].strip() if e.apoyo_visual else ""
        if not titulo or titulo == "Documento completo":  # documento sin encabezados
            titulo = f"Parte {e.orden}"
        y = _encabezado(slide, f"Escena {e.orden} de {total} · {_reloj(inicio)} – {_reloj(fin)}", titulo,
                        alto_titulo=0.75)
        # la duración, como dato destacado a la izquierda; la narración, a la derecha
        _caja(slide, MARGEN_IN, y, 2.6, 2.1, CORAL_SUAVE)
        _texto(slide, MARGEN_IN, y + 0.35, 2.6, 0.9, f"{e.duracion_seg}s", 44, CORAL, negrita=True,
               alinear=PP_ALIGN.CENTER, ajustar=False)
        _texto(slide, MARGEN_IN, y + 1.3, 2.6, 0.4, "de narración", 14, TEXTO_SUAVE, alinear=PP_ALIGN.CENTER,
               ajustar=False)
        x = MARGEN_IN + 2.6 + 0.5
        _texto(slide, x, y, ANCHO_IN - MARGEN_IN - x, ALTO_IN - y - 0.8, e.narracion, 26, TEXTO, minimo=13)
        _pie(slide)
        _notas(slide, f"Narración ({e.duracion_seg} s):\n{e.narracion}")
        inicio = fin


_ETIQUETAS = {
    "Flashcards": "Flashcards",
    "Quiz": "Quiz",
    "Tutorial": "Tutorial",
    "Resumen Ejecutivo": "Resumen ejecutivo",
    "Guion de Clase": "Guion de clase",
}


def generar_presentacion(
    contenido: ContenidoAdaptado,
    titulo: str,
    perfil: str = "",
    score_fidelidad: float | None = None,
) -> bytes:
    """Devuelve el .pptx (bytes) con portada, las diapositivas del formato y cierre."""
    deck = _Deck()
    formato = contenido.formato
    etiqueta = _ETIQUETAS.get(formato, formato)
    if perfil:
        etiqueta = f"{etiqueta} · {perfil}"

    if isinstance(contenido, FlashcardsContenido):
        subtitulo = contenido.introduccion_contextualizada
    elif isinstance(contenido, QuizContenido):
        subtitulo = f"{len(contenido.preguntas)} preguntas para comprobar lo aprendido."
    elif isinstance(contenido, TutorialContenido):
        subtitulo = f"{len(contenido.pasos)} pasos · {contenido.objetivo}"
    elif isinstance(contenido, ResumenEjecutivoContenido):
        subtitulo = "Lo esencial del documento para tomar decisiones."
    else:
        subtitulo = f"{len(contenido.escenas)} escenas · aprox. {contenido.duracion_total_min} min"
    _portada(deck, etiqueta, titulo, subtitulo)

    {
        "Flashcards": _flashcards,
        "Quiz": _quiz,
        "Tutorial": _tutorial,
        "Resumen Ejecutivo": _resumen,
        "Guion de Clase": _guion,
    }[formato](deck, contenido)

    lineas = [f"Material generado por Kairos a partir de «{titulo}»."]
    if score_fidelidad is not None:
        lineas.append(f"Fidelidad a la fuente: {round(score_fidelidad * 100)}%.")
    _cierre(deck, lineas)
    return deck.guardar()

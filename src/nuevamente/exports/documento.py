"""Documento Word (.docx) a partir de cualquier contenido adaptado.

Responsable en el equipo Kairos G10: Dario Higuera Moreno (No Code Developer).

Estructura: título del documento de origen → ficha (formato, perfil, fidelidad) →
el material, con el título de cada sección del documento sobre sus partes. El quiz
lleva las respuestas al final, para poder imprimirlo y resolverlo sin verlas.

Se usa Arial, igual que en la presentación: viene con Office y se ve igual en
Word, Pages y LibreOffice.
"""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

from nuevamente.schemas.formatos import (
    ContenidoAdaptado,
    FlashcardsContenido,
    GuionDeClaseContenido,
    QuizContenido,
    ResumenEjecutivoContenido,
    TutorialContenido,
)

FUENTE = "Arial"
PRIMARIO = RGBColor(0x6C, 0x4E, 0xF5)
CORAL = RGBColor(0xE0, 0x5A, 0x3A)
TEXTO = RGBColor(0x2D, 0x24, 0x40)
TEXTO_SUAVE = RGBColor(0x6E, 0x64, 0x80)

_NUMERACION = re.compile(r"^\d{1,2}(?:\.\d{1,2})*[.)]?\s+")
# caracteres de control que XML no admite (un PDF puede traerlos) y romperían el .docx
_NO_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _limpio(texto: str) -> str:
    return _NO_XML.sub("", texto or "")


class _Doc:
    """Documento con los estilos de Kairos ya aplicados."""

    def __init__(self, titulo: str) -> None:
        self.doc = Document()
        self.doc.core_properties.title = titulo
        self.doc.core_properties.author = "Kairos"
        estilos = self.doc.styles
        estilos["Normal"].font.name = FUENTE
        estilos["Normal"].font.size = Pt(11)
        estilos["Normal"].font.color.rgb = TEXTO
        for nombre, pt, color in (("Title", 24, TEXTO), ("Heading 1", 16, PRIMARIO), ("Heading 2", 12.5, TEXTO)):
            fuente = estilos[nombre].font
            fuente.name, fuente.size, fuente.bold = FUENTE, Pt(pt), True
            fuente.color.rgb = color
        self._seccion = ""

    def titulo(self, texto: str) -> None:
        self.doc.add_heading(_limpio(texto), level=0)

    def apartado(self, texto: str) -> None:
        self.doc.add_heading(_limpio(texto), level=1)

    def subtitulo(self, texto: str) -> None:
        self.doc.add_heading(_limpio(texto), level=2)

    def parrafo(self, texto: str, *, etiqueta: str = "", cursiva=False, suave=False):
        """Párrafo de texto, con una etiqueta opcional en negrita delante ("Narración: ...")."""
        p = self.doc.add_paragraph()
        if etiqueta:
            p.add_run(f"{etiqueta} ").bold = True
        run = p.add_run(_limpio(texto))
        run.italic = cursiva
        if suave:
            run.font.color.rgb = TEXTO_SUAVE
        return p

    def lista(self, textos: list[str], estilo: str = "List Bullet") -> None:
        for texto in textos:
            self.doc.add_paragraph(_limpio(texto), style=estilo)

    def seccion_de(self, parte) -> None:
        """Escribe el título de la sección del documento cuando la parte empieza una nueva."""
        seccion = _NUMERACION.sub("", getattr(parte, "seccion", ""))
        if seccion and seccion != "Documento completo" and seccion != self._seccion:
            self.apartado(seccion)
        self._seccion = seccion

    def salto_de_pagina(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        self._seccion = ""

    def guardar(self) -> bytes:
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()


def _flashcards(d: _Doc, c: FlashcardsContenido) -> None:
    d.parrafo(c.introduccion_contextualizada, suave=True)
    for i, item in enumerate(c.items, 1):
        d.seccion_de(item)
        d.subtitulo(f"Tarjeta {i}. {item.frente}")
        d.parrafo(item.dorso)
        if item.pista_didactica:
            d.parrafo(item.pista_didactica, etiqueta="Pista:", cursiva=True, suave=True)


def _quiz(d: _Doc, c: QuizContenido) -> None:
    for i, p in enumerate(c.preguntas, 1):
        d.seccion_de(p)
        d.subtitulo(f"Pregunta {i}. {p.enunciado}")
        for j, opcion in enumerate(p.opciones):
            d.parrafo(opcion, etiqueta=f"{chr(65 + j)})")
    d.salto_de_pagina()
    d.apartado("Respuestas")
    for i, p in enumerate(c.preguntas, 1):
        d.parrafo(p.opciones[p.indice_correcto], etiqueta=f"{i}. {chr(65 + p.indice_correcto)})")
        d.parrafo(p.justificacion, cursiva=True, suave=True)


def _tutorial(d: _Doc, c: TutorialContenido) -> None:
    d.parrafo(c.objetivo, etiqueta="Objetivo:")
    if c.prerrequisitos:
        d.apartado("Antes de empezar")
        d.lista(c.prerrequisitos)
        d._seccion = ""
    for paso in c.pasos:
        d.seccion_de(paso)
        # si el paso se titula igual que su sección, el título ya está encima
        propio = paso.titulo and _NUMERACION.sub("", paso.titulo) != _NUMERACION.sub("", paso.seccion)
        d.subtitulo(f"Paso {paso.orden}. {paso.titulo}" if propio else f"Paso {paso.orden}")
        d.parrafo(paso.instruccion)
        if paso.resultado_esperado:
            d.parrafo(paso.resultado_esperado, etiqueta="Resultado esperado:", cursiva=True, suave=True)
    if c.errores_comunes:
        d.apartado("Errores comunes")
        d.lista(c.errores_comunes)
    if c.checklist_final:
        d.apartado("Checklist final")
        d.lista([f"☐ {x}" for x in c.checklist_final], estilo="Normal")
    if c.reto_practico:
        d.apartado("Reto práctico")
        d.parrafo(c.reto_practico)


def _resumen(d: _Doc, c: ResumenEjecutivoContenido) -> None:
    d.apartado("Resumen")
    d.parrafo(c.resumen)
    d.apartado("Puntos clave")
    d.lista([_NUMERACION.sub("", pk) for pk in c.puntos_clave])
    if c.decisiones_o_riesgos:
        d.apartado("Riesgos y decisiones")
        d.lista(c.decisiones_o_riesgos)
    if c.impacto_de_negocio:
        d.apartado("Impacto")
        d.parrafo(c.impacto_de_negocio)


def _guion(d: _Doc, c: GuionDeClaseContenido) -> None:
    d.parrafo(
        f"{len(c.escenas)} escenas · duración total aproximada: {c.duracion_total_min} min", suave=True
    )
    inicio = 0
    for escena in c.escenas:
        d.seccion_de(escena)
        d.subtitulo(f"Escena {escena.orden} · {inicio // 60}:{inicio % 60:02d} · {escena.duracion_seg} s")
        d.parrafo(escena.narracion, etiqueta="Narración:")
        if escena.apoyo_visual:
            d.parrafo(escena.apoyo_visual, etiqueta="Apoyo visual:", cursiva=True, suave=True)
        inicio += escena.duracion_seg


_POR_FORMATO = {
    FlashcardsContenido: ("Tarjetas de estudio", _flashcards),
    QuizContenido: ("Quiz", _quiz),
    TutorialContenido: ("Tutorial paso a paso", _tutorial),
    ResumenEjecutivoContenido: ("Resumen ejecutivo", _resumen),
    GuionDeClaseContenido: ("Guion de clase", _guion),
}


def generar_documento(
    contenido: ContenidoAdaptado,
    titulo: str,
    perfil: str = "",
    score_fidelidad: float | None = None,
) -> bytes:
    """Devuelve el .docx del material. `titulo` es el del documento de origen."""
    tipo, escribir = next(
        (v for clase, v in _POR_FORMATO.items() if isinstance(contenido, clase)), (None, None)
    )
    if escribir is None:
        raise ValueError(f"El formato {getattr(contenido, 'formato', '?')} no se exporta a Word")

    d = _Doc(titulo)
    d.titulo(titulo)
    ficha = [tipo, f"Para: {perfil}" if perfil else ""]
    if score_fidelidad is not None:
        ficha.append(f"Fidelidad a la fuente: {round(score_fidelidad * 100)}%")
    etiqueta = d.parrafo(" · ".join(x for x in ficha if x))
    etiqueta.runs[-1].bold = True
    etiqueta.runs[-1].font.color.rgb = CORAL

    escribir(d, contenido)

    pie = d.doc.sections[0].footer.paragraphs[0]
    run = pie.add_run("Generado por Kairos solo a partir del documento proporcionado.")
    run.font.size, run.font.color.rgb = Pt(8), TEXTO_SUAVE
    return d.guardar()

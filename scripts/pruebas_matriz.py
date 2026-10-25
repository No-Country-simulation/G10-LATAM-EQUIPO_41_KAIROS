#!/usr/bin/env python3
"""Matriz de pruebas: un tipo de documento en una categoria, en todas sus posibilidades.

Categoria (nicho) fija: Salud.
Matriz: 4 tipos de archivo (PDF, DOCX, MD, TXT) x 6 formatos x 3 perfiles = 72 pruebas.

Por cada caso se guarda:
  - <NN>_<tipo>_<formato>_<perfil>.json    respuesta completa de la API + detalle interno
  - <NN>_<tipo>_<formato>_<perfil>.png     captura del resultado (render con Pillow)
  - <NN>_<tipo>_<formato>_<perfil>_web.png captura real del navegador (modo --capturas-web)
Y al final:
  - reporte.csv y REPORTE.md               resumen de las 72 pruebas

Uso:
    python scripts/pruebas_matriz.py                 # genera los 72 casos (Gemini real)
    python scripts/pruebas_matriz.py --limite 1      # solo el primer caso (para validar)
    python scripts/pruebas_matriz.py --forzar        # rehace los que ya existan
    python scripts/pruebas_matriz.py --espera 4       # pausa entre casos (cuota)
    python scripts/pruebas_matriz.py --capturas-web   # screenshots reales de la interfaz
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))

FUENTES_DIR = RAIZ / "docs" / "pruebas" / "fuentes"
EVIDENCIA_DIR = RAIZ / "docs" / "pruebas" / "evidencia"
DOCUMENTO_MD = RAIZ / "docs" / "demo" / "protocolo_bioseguridad_salud.md"

NICHO = "Salud"
NIVEL = "Didáctico"

TIPOS = ["pdf", "docx", "md", "txt"]
FORMATOS = [
    "Flashcards",
    "Quiz",
    "Tutorial",
    "Resumen Ejecutivo",
    "Guion de Clase",
    "Podcast",
]
PERFILES = [
    "Principiante",
    "Líder Técnico/Arquitecto",
    "Gestor/Ejecutivo",
]

CONTENT_TYPE = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "md": "text/markdown",
    "txt": "text/plain",
}


def _slug(texto: str) -> str:
    reemplazos = {"í": "i", "á": "a", "é": "e", "ó": "o", "ú": "u", "ñ": "n", "Á": "a"}
    for malo, bueno in reemplazos.items():
        texto = texto.replace(malo, bueno)
    return "".join(c if c.isalnum() else "_" for c in texto.lower()).strip("_")


def _texto_plano(md: str) -> str:
    """Markdown -> texto plano: sin '#' de titulos ni el '>' de las citas."""
    lineas = []
    for linea in md.splitlines():
        limpia = linea.lstrip()
        if limpia.startswith("#"):
            limpia = limpia.lstrip("#").strip()
        elif limpia.startswith(">"):
            limpia = limpia.lstrip(">").strip()
        lineas.append(limpia)
    return "\n".join(lineas).strip() + "\n"


def _secciones_md(md: str) -> list[tuple[int, str]]:
    """Devuelve (nivel, texto) de cada bloque: titulos y parrafos."""
    bloques: list[tuple[int, str]] = []
    parrafo: list[str] = []

    def cerrar():
        if parrafo:
            bloques.append((0, " ".join(parrafo).strip()))
            parrafo.clear()

    for linea in md.splitlines():
        limpia = linea.strip()
        if not limpia:
            cerrar()
        elif limpia.startswith("#"):
            cerrar()
            nivel = len(limpia) - len(limpia.lstrip("#"))
            bloques.append((nivel, limpia.lstrip("#").strip()))
        elif limpia.startswith(">"):
            cerrar()
            bloques.append((0, limpia.lstrip(">").strip()))
        else:
            parrafo.append(limpia)
    cerrar()
    return [b for b in bloques if b[1]]


def crear_fuentes() -> dict[str, Path]:
    """Crea el mismo documento en los 4 formatos soportados."""
    FUENTES_DIR.mkdir(parents=True, exist_ok=True)
    md = DOCUMENTO_MD.read_text(encoding="utf-8")
    rutas: dict[str, Path] = {}

    ruta_md = FUENTES_DIR / "protocolo_salud.md"
    ruta_md.write_text(md, encoding="utf-8")
    rutas["md"] = ruta_md

    ruta_txt = FUENTES_DIR / "protocolo_salud.txt"
    ruta_txt.write_text(_texto_plano(md), encoding="utf-8")
    rutas["txt"] = ruta_txt

    from docx import Document

    doc = Document()
    doc.add_heading("Protocolo General de Bioseguridad para Personal de Salud", level=0)
    for nivel, texto in _secciones_md(md):
        if nivel == 0:
            doc.add_paragraph(texto)
        else:
            doc.add_heading(texto, level=min(nivel, 4))
    ruta_docx = FUENTES_DIR / "protocolo_salud.docx"
    doc.save(str(ruta_docx))
    rutas["docx"] = ruta_docx

    rutas["pdf"] = _crear_pdf(md)
    return rutas


def _crear_pdf(md: str) -> Path:
    """PDF con fpdf2 (fuente base latin-1, igual que el respaldo del proyecto)."""
    from fpdf import FPDF

    reemplazos = {
        "\u2014": "-", "\u2013": "-", "\u201c": '"', "\u201d": '"',
        "\u2018": "'", "\u2019": "'", "\u2026": "...", "\u00a0": " ",
    }

    def limpiar(texto: str) -> str:
        for malo, bueno in reemplazos.items():
            texto = texto.replace(malo, bueno)
        return texto.encode("latin-1", "replace").decode("latin-1")

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    for nivel, texto in _secciones_md(md):
        if nivel:
            pdf.set_font("Helvetica", "B", 15 if nivel <= 1 else 13)
            pdf.multi_cell(0, 8, limpiar(texto))
        else:
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, 6, limpiar(texto))
        pdf.ln(2)
    ruta = FUENTES_DIR / "protocolo_salud.pdf"
    pdf.output(str(ruta))
    return ruta


# --------------------------------------------------------------------------- #
# Captura PNG con Pillow
# --------------------------------------------------------------------------- #

def _fuente(tamano: int, negrita: bool = False):
    from PIL import ImageFont

    candidatos = [
        Path("C:/Windows/Fonts/arialbd.ttf" if negrita else "C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if negrita else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for ruta in candidatos:
        if ruta.exists():
            return ImageFont.truetype(str(ruta), tamano)
    return ImageFont.load_default()


def _envolver(texto: str, font, ancho: int, draw) -> list[str]:
    lineas: list[str] = []
    for parrafo in texto.split("\n"):
        if not parrafo.strip():
            lineas.append("")
            continue
        actual = ""
        for palabra in parrafo.split(" "):
            prueba = f"{actual} {palabra}".strip()
            if draw.textlength(prueba, font=font) <= ancho:
                actual = prueba
            else:
                if actual:
                    lineas.append(actual)
                actual = palabra
        if actual:
            lineas.append(actual)
    return lineas


def generar_png(datos: dict, destino: Path) -> None:
    from PIL import Image, ImageDraw

    destino.parent.mkdir(parents=True, exist_ok=True)
    caso = datos["caso"]
    resp = datos.get("respuesta") or {}
    detalle = datos.get("detalle") or {}
    ev = (detalle.get("evaluacion_calidad") or resp.get("evaluacion_calidad") or {})

    ancho, alto = 1280, 940
    img = Image.new("RGB", (ancho, alto), "#f4f6fb")
    draw = ImageDraw.Draw(img)

    f_marca = _fuente(26, True)
    f_titulo = _fuente(30, True)
    f_sub = _fuente(19)
    f_eti = _fuente(17, True)
    f_dato = _fuente(17)
    f_mono = _fuente(15)
    f_pie = _fuente(14)

    # Cabecera
    draw.rectangle([0, 0, ancho, 90], fill="#1f6feb")
    draw.text((30, 18), "Kairos · Matriz de pruebas", font=f_marca, fill="white")
    draw.text((30, 52), f"Caso #{caso['n']:02d} / 72 · nicho {NICHO}", font=f_sub, fill="#dbe9ff")

    # Titulo del caso
    draw.text((30, 110), f"{caso['tipo'].upper()}  ·  {caso['formato']}", font=f_titulo, fill="#11213b")
    draw.text((30, 152), f"Perfil: {caso['perfil']}  ·  Nivel: {NIVEL}", font=f_sub, fill="#4b5c78")

    # Panel de estado
    aprobado = ev.get("aprobado_por_critico")
    score = ev.get("anclaje_fuente_score")
    umbral = ev.get("umbral_aplicado")
    modelo = (detalle.get("metadatos") or resp.get("metadatos") or {}).get("modelo_llm", "?")
    status = datos.get("status_code")
    color_estado = "#1a7f37" if status == 200 and aprobado else "#b42318"
    draw.rounded_rectangle([30, 195, ancho - 30, 300], radius=12, fill="white", outline="#d7deea", width=2)
    draw.text((50, 212), "HTTP", font=f_eti, fill="#4b5c78")
    draw.text((50, 236), str(status), font=_fuente(22, True), fill=color_estado)
    draw.text((190, 212), "Anclaje a la fuente", font=f_eti, fill="#4b5c78")
    draw.text((190, 236), "—" if score is None else f"{score:.2f}", font=_fuente(22, True), fill="#11213b")
    draw.text((400, 212), "Umbral aplicado", font=f_eti, fill="#4b5c78")
    draw.text((400, 236), "—" if umbral is None else f"{umbral:.2f}", font=f_dato, fill="#11213b")
    draw.text((560, 212), "Aprobado por Critico", font=f_eti, fill="#4b5c78")
    draw.text((560, 236), "SI" if aprobado else ("NO" if aprobado is False else "—"), font=f_dato, fill=color_estado)
    draw.text((760, 212), "Modelo LLM", font=f_eti, fill="#4b5c78")
    draw.text((760, 236), str(modelo), font=f_dato, fill="#11213b")
    draw.text((50, 268), f"Tiempo: {datos.get('tiempo_segundos', '—')} s   ·   Archivo: {caso['archivo']}   ·   Fecha: {datos.get('fecha', '')}", font=f_pie, fill="#4b5c78")

    # Observaciones
    y = 316
    observaciones = ev.get("observaciones") or ""
    draw.text((30, y), "Observaciones del Critico:", font=f_eti, fill="#4b5c78")
    y += 26
    for linea in _envolver(observaciones, f_dato, ancho - 60, draw)[:3]:
        draw.text((30, y), linea, font=f_dato, fill="#11213b")
        y += 23

    # Contenido generado
    y += 12
    draw.text((30, y), "Contenido generado (extracto):", font=f_eti, fill="#4b5c78")
    y += 26
    contenido = resp.get("contenido_adaptado") or detalle.get("contenido_adaptado") or {}
    texto = json.dumps(contenido, ensure_ascii=False, indent=2)
    disponible = alto - y - 40
    for linea in _envolver(texto, f_mono, ancho - 60, draw):
        if y > alto - 40:
            draw.text((30, y - 24), "…(ver el JSON completo)", font=f_pie, fill="#8a97ab")
            break
        draw.text((30, y), linea, font=f_mono, fill="#243b63")
        y += 19

    draw.line([(0, alto - 34), (ancho, alto - 34)], fill="#d7deea", width=2)
    draw.text((30, alto - 26), f"Evidencia: {destino.name}", font=f_pie, fill="#8a97ab")

    img.save(destino)


# --------------------------------------------------------------------------- #
# Generacion de la matriz
# --------------------------------------------------------------------------- #

def construir_casos() -> list[dict]:
    casos = []
    n = 0
    for tipo in TIPOS:
        for formato in FORMATOS:
            for perfil in PERFILES:
                n += 1
                base = f"{n:02d}_{tipo}_{_slug(formato)}_{_slug(perfil)}"
                casos.append({
                    "n": n,
                    "tipo": tipo,
                    "formato": formato,
                    "perfil": perfil,
                    "nicho": NICHO,
                    "nivel": NIVEL,
                    "base": base,
                    "archivo": f"protocolo_salud.{tipo}",
                })
    return casos


def generar_matriz(limite: int | None, forzar: bool, espera: float, reintentar_no_aprobadas: bool = False) -> int:
    from fastapi.testclient import TestClient

    from nuevamente.api.app import app

    EVIDENCIA_DIR.mkdir(parents=True, exist_ok=True)
    rutas = crear_fuentes()
    casos = construir_casos()
    if limite:
        casos = casos[:limite]

    client = TestClient(app)
    filas = []
    print(f"Matriz: {len(casos)} prueba(s) · nicho {NICHO} · generador segun .env")
    print(f"Fuentes: {FUENTES_DIR}")
    print(f"Evidencia: {EVIDENCIA_DIR}\n")

    for caso in casos:
        json_path = EVIDENCIA_DIR / f"{caso['base']}.json"
        if json_path.exists() and not forzar:
            datos = json.loads(json_path.read_text(encoding="utf-8"))
            met = (datos.get("detalle") or {}).get("metadatos") or {}
            ev = (datos.get("detalle") or {}).get("evaluacion_calidad") or {}
            modelo = str(met.get("modelo_llm") or "")
            salio_de_respaldo = "respaldo" in modelo.lower() or not modelo
            no_aprobada = reintentar_no_aprobadas and not ev.get("aprobado_por_critico")
            if datos.get("status_code") == 200 and not salio_de_respaldo and not no_aprobada:
                print(f"[{caso['n']:02d}/{len(casos)}] {caso['tipo']}·{caso['formato']}·{caso['perfil']}  (ya existe)")
                filas.append(_fila_reporte(caso, datos))
                continue
            motivo = ("no aprobada" if no_aprobada
                      else "antes falló" if datos.get("status_code") != 200
                      else f"usó respaldo ({modelo})")
            print(f"[{caso['n']:02d}/{len(casos)}] {caso['tipo']}·{caso['formato']}·{caso['perfil']}  (reintenta: {motivo})")

        fuente = rutas[caso["tipo"]]
        contenido_bytes = fuente.read_bytes()
        files = {"archivo": (fuente.name, contenido_bytes, CONTENT_TYPE[caso["tipo"]])}
        data = {
            "perfil_destinatario": caso["perfil"],
            "formato_salida": caso["formato"],
            "nicho_sector": NICHO,
            "nivel_detalle": NIVEL,
        }

        t0 = time.perf_counter()
        try:
            r = client.post("/api/v1/adaptar/archivo", files=files, data=data)
        except Exception as exc:  # noqa: BLE001
            r = None
            error = f"{type(exc).__name__}: {exc}"
        duracion = round(time.perf_counter() - t0, 2)

        datos = {
            "caso": {k: caso[k] for k in ("n", "tipo", "formato", "perfil", "nicho", "nivel", "archivo")},
            "fecha": datetime.now().isoformat(timespec="seconds"),
            "tiempo_segundos": duracion,
            "status_code": getattr(r, "status_code", None),
            "error": None if r is not None else error,
        }
        if r is not None and r.status_code == 200:
            publica = r.json()
            datos["respuesta"] = publica
            objeto_id = publica["almacenamiento_oci"]["objeto_id"]
            try:
                detalle = client.get(f"/api/v1/contenidos/{objeto_id}/detalle").json()
                datos["detalle"] = detalle
            except Exception:  # noqa: BLE001
                datos["detalle"] = None
        else:
            datos["respuesta"] = None
            try:
                datos["error"] = r.json()
            except Exception:  # noqa: BLE001
                datos["error"] = getattr(r, "text", error if r is None else "")

        json_path.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
        try:
            generar_png(datos, EVIDENCIA_DIR / f"{caso['base']}.png")
        except Exception as exc:  # noqa: BLE001
            print(f"    (no se pudo generar el PNG: {exc})")

        estado = "OK" if datos["status_code"] == 200 else f"ERROR {datos['status_code']}"
        print(f"[{caso['n']:02d}/{len(casos)}] {caso['tipo']}·{caso['formato']}·{caso['perfil']}  -> {estado} ({duracion}s)")
        filas.append(_fila_reporte(caso, datos))

        if espera and caso is not casos[-1]:
            time.sleep(espera)

    _escribir_reportes(filas)
    return 0


def _fila_reporte(caso: dict, datos: dict) -> dict:
    resp = datos.get("respuesta") or {}
    detalle = datos.get("detalle") or {}
    ev = detalle.get("evaluacion_calidad") or resp.get("evaluacion_calidad") or {}
    met = detalle.get("metadatos") or resp.get("metadatos") or {}
    return {
        "n": caso["n"],
        "tipo_archivo": caso["tipo"],
        "formato": caso["formato"],
        "perfil": caso["perfil"],
        "nicho": caso["nicho"],
        "nivel": caso["nivel"],
        "http": datos.get("status_code"),
        "resultado": "OK" if datos.get("status_code") == 200 else "ERROR",
        "anclaje_score": ev.get("anclaje_fuente_score"),
        "umbral": ev.get("umbral_aplicado"),
        "aprobado": ev.get("aprobado_por_critico"),
        "claridad": ev.get("claridad_pedagogica"),
        "modelo_llm": met.get("modelo_llm"),
        "tiempo_segundos": datos.get("tiempo_segundos"),
        "conceptos": "; ".join(met.get("conceptos_clave") or []),
        "archivo_json": f"{caso['base']}.json",
        "captura_png": f"{caso['base']}.png",
        "captura_web": f"{caso['base']}_web.png",
    }


def _escribir_reportes(filas: list[dict]) -> None:
    campos = list(filas[0].keys())
    csv_path = EVIDENCIA_DIR / "reporte.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=campos)
        writer.writeheader()
        writer.writerows(filas)

    total = len(filas)
    ok = sum(1 for f in filas if f["resultado"] == "OK")
    aprob = sum(1 for f in filas if f["aprobado"])
    score = [f["anclaje_score"] for f in filas if isinstance(f["anclaje_score"], (int, float))]
    media = sum(score) / len(score) if score else 0
    tiempos = [f["tiempo_segundos"] for f in filas if isinstance(f["tiempo_segundos"], (int, float))]
    media_t = sum(tiempos) / len(tiempos) if tiempos else 0

    lineas = [
        "# Reporte de la matriz de pruebas",
        "",
        f"- Categoria (nicho): **{NICHO}**",
        f"- Matriz: **4 tipos de archivo × 6 formatos × 3 perfiles = {total} pruebas**",
        f"- Generado: {datetime.now().isoformat(timespec='seconds')}",
        "",
        "## Resumen",
        "",
        "| Métrica | Valor |",
        "|---|---|",
        f"| Pruebas ejecutadas | {total} |",
        f"| Respuestas HTTP 200 | {ok} |",
        f"| Aprobadas por el Crítico | {aprob} |",
        f"| Anclaje medio | {media:.2f} |",
        f"| Tiempo medio por prueba | {media_t:.1f} s |",
        "",
        "## Detalle por prueba",
        "",
        "| # | Tipo | Formato | Perfil | HTTP | Score | Aprobado | Modelo | s | JSON |",
        "|---:|---|---|---|---:|---:|:--:|---|---:|---|",
    ]
    for f in filas:
        score_txt = f"{f['anclaje_score']:.2f}" if isinstance(f["anclaje_score"], (int, float)) else "—"
        aprob_txt = "✅" if f["aprobado"] else ("✗" if f["aprobado"] is False else "—")
        lineas.append(
            f"| {f['n']} | {f['tipo_archivo']} | {f['formato']} | {f['perfil']} | {f['http']} | "
            f"{score_txt} | {aprob_txt} | {f['modelo_llm']} | {f['tiempo_segundos']} | {f['archivo_json']} |"
        )
    lineas += [
        "",
        "## Evidencia por prueba",
        "",
        "Cada prueba tiene: el JSON (respuesta completa), un PNG del resultado y un PNG de la "
        "interfaz web real. Ver también `index.html` (galería con todas las capturas).",
        "",
        "| # | JSON | Captura resultado | Captura interfaz |",
        "|---:|---|---|---|",
    ]
    for f in filas:
        lineas.append(
            f"| {f['n']} | `{f['archivo_json']}` | `{f['captura_png']}` | `{f['captura_web']}` |"
        )
    (EVIDENCIA_DIR / "REPORTE.md").write_text("\n".join(lineas) + "\n", encoding="utf-8")
    _escribir_galeria(filas)

    print(f"\nReporte: {csv_path}")
    print(f"Reporte: {EVIDENCIA_DIR / 'REPORTE.md'}")
    print(f"Resumen: {ok}/{total} HTTP 200 · {aprob} aprobadas · anclaje medio {media:.2f}")


def _escribir_galeria(filas: list[dict]) -> None:
    """Galería HTML con las capturas de las 72 pruebas."""
    total = len(filas)
    ok = sum(1 for f in filas if f["resultado"] == "OK")
    aprob = sum(1 for f in filas if f["aprobado"])
    tarjetas = []
    for f in filas:
        score = f"{f['anclaje_score']:.2f}" if isinstance(f["anclaje_score"], (int, float)) else "—"
        clase = "ok" if f["aprobado"] else "mal"
        tarjetas.append(
            f'<article class="tarjeta {clase}">'
            f'<header><span class="n">#{f["n"]:02d}</span>'
            f'<span class="tag">{f["tipo_archivo"].upper()}</span>'
            f'<span class="tag">{f["formato"]}</span>'
            f'<span class="tag">{f["perfil"]}</span></header>'
            f'<a href="{f["captura_web"]}" target="_blank">'
            f'<img loading="lazy" src="{f["captura_web"]}" alt="Caso {f["n"]}"></a>'
            f'<footer>HTTP {f["http"]} · score {score} · {"aprobada" if f["aprobado"] else "no aprobada"}'
            f' · {f["tiempo_segundos"]} s · <a href="{f["archivo_json"]}" target="_blank">JSON</a>'
            f' · <a href="{f["captura_png"]}" target="_blank">captura</a></footer>'
            f'</article>'
        )
    html = f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Matriz de pruebas · Nicho {NICHO}</title>
<style>
 body{{font-family:system-ui,Segoe UI,Arial,sans-serif;background:#f4f6fb;color:#11213b;margin:0;padding:24px}}
 h1{{margin:0 0 4px}}
 .resumen{{color:#4b5c78;margin:0 0 20px}}
 .grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(360px,1fr));gap:16px}}
 .tarjeta{{background:#fff;border:1px solid #d7deea;border-radius:12px;overflow:hidden;display:flex;flex-direction:column}}
 .tarjeta.ok{{border-left:6px solid #1a7f37}}
 .tarjeta.mal{{border-left:6px solid #b42318}}
 header{{display:flex;flex-wrap:wrap;gap:6px;padding:10px 12px;align-items:center;border-bottom:1px solid #eef2f7}}
 .n{{font-weight:700}}
 .tag{{background:#eaf1ff;color:#1f6feb;border-radius:999px;padding:2px 9px;font-size:12px}}
 img{{width:100%;display:block;background:#fff;max-height:520px;object-fit:cover;object-position:top}}
 footer{{padding:8px 12px;font-size:12px;color:#4b5c78}}
 a{{color:#1f6feb}}
</style></head><body>
<h1>Matriz de pruebas · NuevaMente (Kairos)</h1>
<p class="resumen">Categoría (nicho): <b>{NICHO}</b> · 4 tipos de archivo × 6 formatos × 3 perfiles = <b>{total}</b> pruebas.
 {ok} respuestas HTTP 200 · {aprob} aprobadas por el Crítico. Generado {datetime.now().isoformat(timespec="seconds")}.</p>
<div class="grid">
{chr(10).join(tarjetas)}
</div></body></html>"""
    (EVIDENCIA_DIR / "index.html").write_text(html, encoding="utf-8")
    print(f"Galería: {EVIDENCIA_DIR / 'index.html'}")


# --------------------------------------------------------------------------- #
# Capturas reales del navegador (Playwright + Edge)
# --------------------------------------------------------------------------- #

def capturas_web(puerto: int = 8121) -> int:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("Falta Playwright. Instala con: pip install playwright", file=sys.stderr)
        return 1

    jsons = sorted(EVIDENCIA_DIR.glob("*.json"))
    jsons = [p for p in jsons if p.name != "resumen_72.json"]
    if not jsons:
        print("No hay evidencia JSON. Corre primero la generación.", file=sys.stderr)
        return 1

    base = f"http://127.0.0.1:{puerto}"
    proceso = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "nuevamente.api.app:app", "--port", str(puerto)],
        cwd=str(RAIZ),
    )
    try:
        import urllib.request

        for _ in range(60):
            try:
                urllib.request.urlopen(f"{base}/health", timeout=1)
                break
            except Exception:  # noqa: BLE001
                time.sleep(0.5)

        print(f"Servidor en {base}. Capturando {len(jsons)} materiales…")
        with sync_playwright() as p:
            navegador = p.chromium.launch(channel="msedge", headless=True)
            pagina = navegador.new_page(viewport={"width": 1440, "height": 1600})
            for i, jp in enumerate(jsons, 1):
                datos = json.loads(jp.read_text(encoding="utf-8"))
                resp = datos.get("respuesta") or {}
                objeto_id = (resp.get("almacenamiento_oci") or {}).get("objeto_id")
                if not objeto_id:
                    continue
                titulo = resp.get("contenido_adaptado", {}).get("titulo") or datos["caso"]["archivo"]
                entrada = [{
                    "objeto_id": objeto_id,
                    "titulo": titulo,
                    "formato": datos["caso"]["formato"],
                    "perfil": datos["caso"]["perfil"],
                    "fecha": datos.get("fecha", ""),
                }]
                pagina.goto(base, wait_until="networkidle")
                pagina.evaluate("h => localStorage.setItem('nuevamente.historial', JSON.stringify(h))", entrada)
                pagina.reload(wait_until="networkidle")
                try:
                    pagina.click("#historial li button", timeout=5000)
                    pagina.wait_for_selector("#resultado", state="visible", timeout=15000)
                    time.sleep(0.6)
                    destino = EVIDENCIA_DIR / f"{jp.stem}_web.png"
                    pagina.screenshot(path=str(destino), full_page=True)
                    print(f"[{i}/{len(jsons)}] {jp.stem} -> {destino.name}")
                except Exception as exc:  # noqa: BLE001
                    print(f"[{i}/{len(jsons)}] {jp.stem}: no se pudo capturar ({exc})")
            navegador.close()
    finally:
        proceso.terminate()
        try:
            proceso.wait(timeout=10)
        except Exception:  # noqa: BLE001
            proceso.kill()
    return 0


def reconstruir_reporte() -> int:
    """Rehace reporte.csv, REPORTE.md e index.html a partir de la evidencia ya generada."""
    casos = construir_casos()
    filas = []
    for caso in casos:
        json_path = EVIDENCIA_DIR / f"{caso['base']}.json"
        if not json_path.exists():
            continue
        datos = json.loads(json_path.read_text(encoding="utf-8"))
        filas.append(_fila_reporte(caso, datos))
    if not filas:
        print("No hay evidencia para reportar.", file=sys.stderr)
        return 1
    _escribir_reportes(filas)
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Matriz de pruebas de tipos de documento (nicho Salud).")
    parser.add_argument("--limite", type=int, default=None, help="Procesa solo los primeros N casos.")
    parser.add_argument("--forzar", action="store_true", help="Rehace los casos que ya tienen evidencia.")
    parser.add_argument("--espera", type=float, default=3.0, help="Segundos de pausa entre casos (cuota).")
    parser.add_argument("--capturas-web", action="store_true", help="Toma screenshots reales de la interfaz.")
    parser.add_argument("--puerto", type=int, default=8121, help="Puerto para las capturas web.")
    parser.add_argument("--proveedor", default=None, help="Fuerza LLM_PROVIDER (p. ej. gemini, template, auto).")
    parser.add_argument("--reintentar-no-aprobadas", action="store_true", help="Rehace los casos no aprobados por el Crítico.")
    parser.add_argument("--reporte", action="store_true", help="Solo regenera reporte.csv, REPORTE.md e index.html.")
    args = parser.parse_args()

    if args.proveedor:
        os.environ["LLM_PROVIDER"] = args.proveedor

    if args.reporte:
        return reconstruir_reporte()
    if args.capturas_web:
        return capturas_web(args.puerto)
    return generar_matriz(args.limite, args.forzar, args.espera, args.reintentar_no_aprobadas)


if __name__ == "__main__":
    raise SystemExit(main())
